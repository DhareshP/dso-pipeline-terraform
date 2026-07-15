"""Export utilities for roster data."""
import io
from datetime import date
from typing import Dict, List, Any
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, Paragraph, Spacer, TableStyle


def export_to_excel(
    roster: Dict,
    assignments: List[Dict],
    shifts: List[Dict]
) -> io.BytesIO:
    """
    Export roster to Excel format.

    Args:
        roster: Roster metadata
        assignments: List of assignments with employee and shift details
        shifts: List of shift configurations

    Returns:
        BytesIO buffer containing the Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"

    # Styles
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Shift colors
    shift_colors = {
        "morning": PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
        "afternoon": PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid"),
        "night": PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid"),
        "weekly_off": PatternFill(start_color="D0CECE", end_color="D0CECE", fill_type="solid"),
    }

    # Title
    ws.merge_cells("A1:Z1")
    ws["A1"] = f"Roster: {roster['name']} ({roster['start_date']} to {roster['end_date']})"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = center_align

    # Build date list
    start_date = date.fromisoformat(roster['start_date'])
    end_date = date.fromisoformat(roster['end_date'])
    num_days = (end_date - start_date).days + 1
    dates = [start_date + __import__('datetime').timedelta(days=i) for i in range(num_days)]

    # Build employee list from assignments
    employee_assignments = defaultdict(dict)
    for a in assignments:
        emp_id = a["employee_id"]
        emp_name = a.get("employees", {}).get("name", "Unknown")
        d = a["date"]
        if a["is_weekly_off"]:
            employee_assignments[(emp_id, emp_name)][d] = ("WO", None)
        else:
            shift_id = a.get("shift_id")
            shift_name = a.get("shifts", {}).get("name", "Shift") if a.get("shifts") else "Shift"
            employee_assignments[(emp_id, emp_name)][d] = (shift_name, shift_id)

    # Header row
    ws.cell(row=3, column=1, value="Employee")
    ws.cell(row=3, column=1).font = header_font_white
    ws.cell(row=3, column=1).fill = header_fill
    ws.cell(row=3, column=1).alignment = center_align
    ws.cell(row=3, column=1).border = thin_border

    for col_idx, d in enumerate(dates, start=2):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = d.strftime("%a\n%d/%m")
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = 8

    ws.column_dimensions['A'].width = 20

    # Data rows
    row_idx = 4
    for (emp_id, emp_name), emp_data in sorted(employee_assignments.items(), key=lambda x: x[0][1]):
        # Employee name
        cell = ws.cell(row=row_idx, column=1, value=emp_name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

        # Daily assignments
        for col_idx, d in enumerate(dates, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            d_str = d.isoformat()
            if d_str in emp_data:
                shift_name, shift_id = emp_data[d_str]
                cell.value = shift_name
                if shift_name == "WO":
                    cell.fill = shift_colors["weekly_off"]
                else:
                    # Match shift color
                    s_lower = shift_name.lower() if shift_name else ""
                    if "morning" in s_lower:
                        cell.fill = shift_colors["morning"]
                    elif "afternoon" in s_lower:
                        cell.fill = shift_colors["afternoon"]
                    elif "night" in s_lower:
                        cell.fill = shift_colors["night"]
                    else:
                        cell.fill = shift_colors["weekly_off"]
            cell.alignment = center_align
            cell.border = thin_border

        row_idx += 1

    # Legend
    row_idx += 2
    ws.cell(row=row_idx, column=1, value="Legend:")
    ws.cell(row=row_idx, column=1).font = Font(bold=True)
    row_idx += 1

    legend_items = [
        ("Morning (M)", shift_colors["morning"]),
        ("Afternoon (A)", shift_colors["afternoon"]),
        ("Night (N)", shift_colors["night"]),
        ("Weekly Off (WO)", shift_colors["weekly_off"]),
    ]
    for text, fill in legend_items:
        cell = ws.cell(row=row_idx, column=1, value=text)
        cell.fill = fill
        cell.alignment = center_align
        cell.border = thin_border
        row_idx += 1

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer


def export_to_csv(
    roster: Dict,
    assignments: List[Dict],
    shifts: List[Dict]
) -> io.BytesIO:
    """
    Export roster to CSV format.

    Args:
        roster: Roster metadata
        assignments: List of assignments
        shifts: List of shift configurations

    Returns:
        BytesIO buffer containing the CSV content
    """
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Build date list
    start_date = date.fromisoformat(roster['start_date'])
    end_date = date.fromisoformat(roster['end_date'])
    num_days = (end_date - start_date).days + 1
    dates = [start_date + __import__('datetime').timedelta(days=i) for i in range(num_days)]

    # Build employee assignments
    employee_assignments = defaultdict(dict)
    for a in assignments:
        emp_name = a.get("employees", {}).get("name", "Unknown")
        d = a["date"]
        if a["is_weekly_off"]:
            employee_assignments[emp_name][d] = "WO"
        else:
            shift_name = a.get("shifts", {}).get("name", "Shift") if a.get("shifts") else "Shift"
            employee_assignments[emp_name][d] = shift_name

    # Header
    header = ["Employee"] + [d.strftime("%Y-%m-%d") for d in dates]
    writer.writerow(header)

    # Data rows
    for emp_name in sorted(employee_assignments.keys()):
        row = [emp_name]
        for d in dates:
            row.append(employee_assignments[emp_name].get(d, "WO"))
        writer.writerow(row)

    # Convert string buffer to bytes buffer
    output_buffer = io.BytesIO()
    output_buffer.write(buffer.getvalue().encode('utf-8'))
    buffer.close()
    return output_buffer


def export_to_pdf(
    roster: Dict,
    assignments: List[Dict],
    shifts: List[Dict]
) -> io.BytesIO:
    """
    Export roster to PDF format.

    Args:
        roster: Roster metadata
        assignments: List of assignments
        shifts: List of shift configurations

    Returns:
        BytesIO buffer containing the PDF
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    elements = []
    styles = getSampleStyleSheet()

    # Title
    title = Paragraph(
        f"<b>Roster: {roster['name']}</b><br/>"
        f"Period: {roster['start_date']} to {roster['end_date']}",
        styles['Heading1']
    )
    elements.append(title)
    elements.append(Spacer(1, 20))

    # Build date list
    start_date = date.fromisoformat(roster['start_date'])
    end_date = date.fromisoformat(roster['end_date'])
    num_days = (end_date - start_date).days + 1

    # Limit days per page for readability
    days_per_page = 14
    total_pages = (num_days + days_per_page - 1) // days_per_page

    dates = [start_date + __import__('datetime').timedelta(days=i) for i in range(num_days)]

    # Build employee assignments
    employee_assignments = defaultdict(dict)
    for a in assignments:
        emp_name = a.get("employees", {}).get("name", "Unknown")
        d = a["date"]
        if a["is_weekly_off"]:
            employee_assignments[emp_name][d] = "WO"
        else:
            shift_name = a.get("shifts", {}).get("name", "Shift") if a.get("shifts") else "Shift"
            employee_assignments[emp_name][d] = shift_name

    # Create table data for each page
    for page in range(total_pages):
        start_idx = page * days_per_page
        end_idx = min(start_idx + days_per_page, num_days)
        page_dates = dates[start_idx:end_idx]

        # Header row
        header = ["Employee"] + [d.strftime("%a %d/%m") for d in page_dates]

        # Data rows
        data = [header]
        for emp_name in sorted(employee_assignments.keys()):
            row = [emp_name]
            for d in page_dates:
                shift = employee_assignments[emp_name].get(d, "WO")
                # Abbreviate shift names for PDF
                if "morning" in shift.lower():
                    shift = "M"
                elif "afternoon" in shift.lower():
                    shift = "A"
                elif "night" in shift.lower():
                    shift = "N"
                elif shift == "WO":
                    shift = "WO"
                row.append(shift)
            data.append(row)

        # Create table
        col_widths = [100] + [40] * len(page_dates)
        table = Table(data, colWidths=col_widths)

        # Table style
        table_style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ])

        # Color cells based on shift type
        for row_idx, row in enumerate(data[1:], start=1):
            for col_idx, cell in enumerate(row[1:], start=1):
                if cell == "M":
                    table_style.add('BACKGROUND', (col_idx, row_idx), (col_idx, row_idx), colors.HexColor('#FFC000'))
                elif cell == "A":
                    table_style.add('BACKGROUND', (col_idx, row_idx), (col_idx, row_idx), colors.HexColor('#70AD47'))
                elif cell == "N":
                    table_style.add('BACKGROUND', (col_idx, row_idx), (col_idx, row_idx), colors.HexColor('#5B9BD5'))
                elif cell == "WO":
                    table_style.add('BACKGROUND', (col_idx, row_idx), (col_idx, row_idx), colors.HexColor('#D0CECE'))

        table.setStyle(table_style)
        elements.append(table)

        if page < total_pages - 1:
            elements.append(Spacer(1, 30))

    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    return buffer
